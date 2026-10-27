"""Build the monthly kWh report and its Excel workbook.

The data comes from storage.daily_rows(); this module only shapes it for the
dashboard and for the .xlsx export.
"""

from __future__ import annotations

import calendar
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import storage

_HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_STATUS_LABEL = {"complete": "lengkap", "partial": "sebagian", "missing": "tanpa data"}


def month_bounds(year: int, month: int) -> tuple[date, date]:
    last_day = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last_day)


def _sum(values: list[float | None]) -> float | None:
    numbers = [value for value in values if value is not None]
    return sum(numbers) if numbers else None


def _has_data(value: dict) -> bool:
    return value["kwh"] is not None or value["t1"] is not None or value["t2"] is not None


def _meter_totals(rows: list[dict]) -> dict:
    return {
        "kwh": _sum([row["kwh"] for row in rows]),
        "t1": _sum([row["kwh_t1"] for row in rows]),
        "t2": _sum([row["kwh_t2"] for row in rows]),
    }


def build_report(year: int, month: int, meters: list[dict]) -> dict:
    """Daily kWh for every configured meter over one calendar month."""
    first_day, last_day = month_bounds(year, month)
    stored = {}
    for row in storage.daily_rows(first_day, last_day):
        stored.setdefault(row["day"], {})[row["meter_id"]] = row

    days = []
    partial_days = []
    missing_days = []
    day = first_day
    while day <= last_day:
        per_meter = stored.get(day, {})
        day_meters = {}
        for meter in meters:
            row = per_meter.get(meter["id"])
            day_meters[meter["id"]] = {
                "kwh": row["kwh"] if row else None,
                "t1": row["kwh_t1"] if row else None,
                "t2": row["kwh_t2"] if row else None,
                "complete": bool(row["complete"]) if row else False,
            }

        kwh = _sum([value["kwh"] for value in day_meters.values()])
        t1 = _sum([value["t1"] for value in day_meters.values()])
        t2 = _sum([value["t2"] for value in day_meters.values()])
        t1t2 = (t1 + t2) if (t1 is not None and t2 is not None) else None

        # A meter that is not installed yet reports nothing at all; that must
        # not mark the whole day as partial. Only judge the meters that have
        # readings, and call the day missing when none of them has any.
        with_data = [value for value in day_meters.values() if _has_data(value)]
        if not with_data:
            status = "missing"
            missing_days.append(day.day)
        elif any(not value["complete"] for value in with_data):
            status = "partial"
            partial_days.append(day.day)
        else:
            status = "complete"

        days.append({
            "day": day.isoformat(),
            "dayNumber": day.day,
            "month": day.month,
            "weekday": day.weekday(),
            "meters": day_meters,
            "total": {
                "kwh": kwh,
                "t1": t1,
                "t2": t2,
                "t1t2": t1t2,
                "diff": (kwh - t1t2) if (kwh is not None and t1t2 is not None) else None,
            },
            "status": status,
        })
        day = date.fromordinal(day.toordinal() + 1)

    meter_totals = {}
    for meter in meters:
        rows = [row for row in stored.values() if meter["id"] in row]
        meter_totals[meter["id"]] = _meter_totals([row[meter["id"]] for row in rows])

    month_totals = {
        "kwh": _sum([entry["total"]["kwh"] for entry in days]),
        "t1": _sum([entry["total"]["t1"] for entry in days]),
        "t2": _sum([entry["total"]["t2"] for entry in days]),
    }
    month_totals["t1t2"] = (
        (month_totals["t1"] + month_totals["t2"])
        if (month_totals["t1"] is not None and month_totals["t2"] is not None) else None
    )

    return {
        "year": year,
        "month": month,
        "meters": [{"id": meter["id"], "name": meter["name"]} for meter in meters],
        "days": days,
        "meterTotals": meter_totals,
        "monthTotals": month_totals,
        "partialDays": partial_days,
        "missingDays": missing_days,
        "hasData": any(entry["total"]["kwh"] is not None for entry in days),
    }


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for index in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def build_workbook(report: dict) -> BytesIO:
    workbook = Workbook()
    meters = report["meters"]
    days = report["days"]

    # Sheet 1: one row per day, one column per meter plus the factory total.
    sheet = workbook.active
    sheet.title = "Harian per Mesin"
    _write_header(sheet, ["Tanggal", "Hari"] + [m["id"] for m in meters] + ["TOTAL kWh", "Status"])
    weekday_names = ["Sen", "Sel", "Rab", "Kam", "Jum", "Sab", "Min"]
    for entry in days:
        sheet.append(
            [
                f"{entry['dayNumber']:02d}/{report['month']:02d}/{report['year']}",
                weekday_names[entry["weekday"]],
            ]
            + [entry["meters"][m["id"]]["kwh"] for m in meters]
            + [entry["total"]["kwh"], _STATUS_LABEL[entry["status"]]]
        )
    sheet.append([])
    sheet.append(["TOTAL"] + [""] * len(meters) + [report["monthTotals"]["kwh"], ""])
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=3, max_col=3 + len(meters)):
        for cell in row:
            cell.number_format = "0.0"

    # Sheet 2: the tariff breakdown behind every day.
    detail = workbook.create_sheet("Rincian Tarif")
    _write_header(detail, ["Tanggal", "Mesin", "Nama", "Tarif 1 (kWh)", "Tarif 2 (kWh)",
                           "T1+T2 (kWh)", "Total (kWh)", "Selisih", "Status"])
    for entry in days:
        for meter in meters:
            value = entry["meters"][meter["id"]]
            t1, t2 = value["t1"], value["t2"]
            t1t2 = (t1 + t2) if (t1 is not None and t2 is not None) else None
            total = value["kwh"]
            detail.append([
                entry["day"], meter["id"], meter["name"], t1, t2, t1t2, total,
                (total - t1t2) if (total is not None and t1t2 is not None) else None,
                _STATUS_LABEL[entry["status"]],
            ])
    for row in detail.iter_rows(min_row=2, max_row=detail.max_row, min_col=4, max_col=8):
        for cell in row:
            cell.number_format = "0.0"

    # Sheet 3: one line per meter for the whole month.
    summary = workbook.create_sheet("Ringkasan Bulanan")
    _write_header(summary, ["Mesin", "Nama", "Tarif 1 (kWh)", "Tarif 2 (kWh)", "Total (kWh)"])
    for meter in meters:
        totals = report["meterTotals"][meter["id"]]
        summary.append([meter["id"], meter["name"], totals["t1"], totals["t2"], totals["kwh"]])
    summary.append([])
    summary.append(["TOTAL", "", report["monthTotals"]["t1"], report["monthTotals"]["t2"],
                    report["monthTotals"]["kwh"]])
    for row in summary.iter_rows(min_row=2, max_row=summary.max_row, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = "0.0"

    for ws in workbook.worksheets:
        for index, column in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
            width = max(len(str(cell.value or "")) for cell in column) + 4
            ws.column_dimensions[get_column_letter(index)].width = max(12, min(width, 26))

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _style_header_row(sheet, values: list[str], row: int) -> None:
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=index, value=value)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def build_profile_workbook(profile: dict, heatmap: dict) -> BytesIO:
    """Two sheets: the hourly load profile and the monthly heatmap."""
    workbook = Workbook()
    scope = "Total pabrik" if profile["selection"] == "all" else profile["selection"]

    daily = workbook.active
    daily.title = "Beban per Jam"
    daily.append([f"PROFIL BEBAN — {profile['date']} — {scope}"])
    daily["A1"].font = Font(bold=True, size=12)
    daily.append(["Beban puncak (kW)", profile["peakKw"]])
    daily.append(["Jam terpuncak",
                  f"{profile['peakHour']:02d}:00" if profile["peakHour"] is not None else "-"])
    daily.append(["Total energi (kWh)", round(profile["totalKwh"], 3)])
    daily.append([])

    header_row = daily.max_row + 1
    _style_header_row(daily, ["Jam", "Rata-rata kW", "Energi (kWh)", "Cakupan (menit)", "Status"], header_row)
    for hour in profile["hours"]:
        daily.append([hour["label"], hour["kw"], hour["kwh"], hour["coveredMinutes"],
                      _STATUS_LABEL[hour["status"]]])
    for row in daily.iter_rows(min_row=header_row + 1, max_row=daily.max_row, min_col=2, max_col=3):
        for cell in row:
            cell.number_format = "0.00"
    daily.column_dimensions["A"].width = 8
    for column in "BCDE":
        daily.column_dimensions[column].width = 16
    daily.freeze_panes = f"A{header_row + 1}"

    heat = workbook.create_sheet("Heatmap Bulanan")
    heat.append([f"HEATMAP BEBAN — {calendar.month_name[heatmap['month']]} {heatmap['year']} — {scope}"])
    heat["A1"].font = Font(bold=True, size=12)
    heat.append([])
    header_row = heat.max_row + 1
    _style_header_row(heat, ["Tgl"] + [f"{hour:02d}" for hour in range(24)] + ["Total kWh"], header_row)

    for day in heatmap["days"]:
        heat.append([day["day"]]
                    + [None if value is None else round(value, 2) for value in day["hours"]]
                    + [round(day["totalKwh"], 2)])

    first = header_row + 1
    last = header_row + len(heatmap["days"])
    if last >= first:
        heat.conditional_formatting.add(
            f"B{first}:Y{last}",
            ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                           end_type="num", end_value=max(heatmap["maxKw"], 0.001), end_color="168EC8"),
        )
    for row in heat.iter_rows(min_row=first, max_row=last, min_col=2, max_col=26):
        for cell in row:
            cell.number_format = "0.0"
    heat.column_dimensions["A"].width = 6
    for column in range(2, 26):
        heat.column_dimensions[get_column_letter(column)].width = 5
    heat.column_dimensions["Z"].width = 12
    heat.freeze_panes = f"B{first}"

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
